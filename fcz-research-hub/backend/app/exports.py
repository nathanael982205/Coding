"""Exporte (CSV, Excel, JSON, Word, BibTeX/RIS) sowie Backup und Wiederherstellung."""
from __future__ import annotations

import csv
import io
import json
import re
import shutil
import sqlite3
import tempfile
import zipfile
from pathlib import Path
from typing import Any

from . import config, db, services
from .schema import (CHAPTERS, DESIGN_CATEGORIES, LEVEL_CATEGORIES, RECOMMENDATIONS, STAT_METHODS,
                     TOPICS, vocab)

MATRIX_COLUMNS = [
    ("id", "ID"), ("ref", "Kurzbeleg"), ("title", "Titel"), ("authors", "Autoren"), ("year", "Jahr"),
    ("journal", "Journal"), ("doi", "DOI"), ("design", "Forschungsdesign"), ("study_type", "Studientyp (Autoren)"),
    ("research_question", "Forschungsfrage"), ("n", "Stichprobe (n)"), ("sample_desc", "Teilnehmer"),
    ("level", "Leistungsniveau"), ("sport", "Sportart"), ("gps_system", "GPS-System"),
    ("gps_variables", "GPS-Variablen"), ("cmj_device", "CMJ-Gerät"), ("cmj_variables", "CMJ-Parameter"),
    ("timepoints", "Messzeitpunkte"), ("md_offsets", "MD-Bezug"), ("statistics", "Statistik"),
    ("stat_tags", "Statistik-Kategorien"), ("key_findings", "Wichtigste Ergebnisse"),
    ("limitations", "Limitationen (Autoren)"), ("relevance", "Relevanz-Score"),
    ("recommendations", "Empfehlung"), ("chapters", "Theoriekapitel"), ("topics", "Themen"),
    ("user_tags", "Markierungen"), ("status", "Status"), ("filename", "Datei"),
]


def _cv(node: Any, with_page: bool = True) -> str:
    if not isinstance(node, dict):
        return ""
    v = str(node.get("value", ""))
    if with_page and node.get("page") and node.get("status") != "not_reported":
        lbl = f", gedr. S. {node['page_label']}" if node.get("page_label") else ""
        v += f" [PDF-S. {node['page']}{lbl}]"
    if node.get("status") == "uncertain":
        v += " (unsicher)"
    return v


def _list(nodes: Any) -> str:
    return " | ".join(_cv(n) for n in (nodes or []) if isinstance(n, dict) and n.get("status") != "not_reported")


def matrix_rows() -> list[dict[str, Any]]:
    tags = vocab()["user_tags"]
    chapter_titles = {c["key"]: c["title"] for c in CHAPTERS}
    out = []
    for s in services.list_studies():
        row = services.get_study_row(s["id"])
        a = db.loads(row.get("analysis_json"), {}) or {}
        d, sm, m, r = a.get("design", {}), a.get("sample", {}), a.get("methods", {}), a.get("results", {})
        out.append({
            "id": s["id"], "ref": services.short_ref(row), "title": s["title"] or s["filename"],
            "authors": s["authors"], "year": s["year"], "journal": s["journal"], "doi": s["doi"],
            "design": DESIGN_CATEGORIES.get(s["design_category"] or "", ""),
            "study_type": _cv(d.get("study_type")), "research_question": _cv(d.get("research_question")),
            "n": s["sample_n"], "sample_desc": _cv(sm.get("n_participants")),
            "level": LEVEL_CATEGORIES.get(s["level_category"] or "", ""), "sport": s["sport"],
            "gps_system": _cv(m.get("gps_system")), "gps_variables": "; ".join(s["gps_variables"]),
            "cmj_device": _cv(m.get("cmj_device")), "cmj_variables": "; ".join(s["cmj_variables"]),
            "timepoints": _cv(m.get("measurement_timing")),
            "md_offsets": ", ".join((a.get("timepoints") or {}).get("md_offsets") or []),
            "statistics": _cv(m.get("statistical_methods")),
            "stat_tags": ", ".join(STAT_METHODS.get(t, t) for t in s["stat_methods"]),
            "key_findings": _list(r.get("key_findings")),
            "limitations": _list((a.get("critical") or {}).get("limitations_reported")),
            "relevance": s["relevance_score"],
            "recommendations": "; ".join(RECOMMENDATIONS.get(x, x) for x in s["recommendations"]),
            "chapters": "; ".join(chapter_titles.get(x, x) for x in s["chapters"]),
            "topics": ", ".join(TOPICS.get(t, t) for t in s["topics"]),
            "user_tags": ", ".join(tags.get(t, t) for t in s["user_tags"]),
            "status": s["status"], "filename": s["filename"],
        })
    return out


def to_csv() -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    w.writerow([label for _, label in MATRIX_COLUMNS])
    for r in matrix_rows():
        w.writerow(["" if r.get(k) is None else r.get(k) for k, _ in MATRIX_COLUMNS])
    return ("﻿" + buf.getvalue()).encode("utf-8")  # BOM für Excel


def to_xlsx() -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Literaturmatrix"
    ws.append([label for _, label in MATRIX_COLUMNS])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1E3A6E")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in matrix_rows():
        ws.append([r.get(k) for k, _ in MATRIX_COLUMNS])
    widths = {"title": 50, "research_question": 50, "key_findings": 70, "limitations": 50, "sample_desc": 35,
              "statistics": 40, "timepoints": 35}
    for i, (k, _) in enumerate(MATRIX_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(k, 18)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions

    ev = wb.create_sheet("Belege")
    ev.append(["Studie-ID", "Feld", "Aussage", "PDF-Seite", "Gedruckte Seite", "Originalzitat", "Status",
               "Quelle", "Zitatprüfung"])
    for c in ev[1]:
        c.font = Font(bold=True)
    for e in services.evidence_for_study_ids([s["id"] for s in services.list_studies()]):
        ev.append([e["study_id"], e["field_path"], e["statement"], e["page"], e["page_label"], e["quote"],
                   e["status"], e["source_type"], e["quote_check"]])
    info = wb.create_sheet("Hinweise")
    for line in [
        "FCZ Thesis Research Hub – Export der Literaturmatrix",
        f"Erstellt: {db.now_iso()}",
        "PDF-S. = physische Seite im PDF; gedr. S. = gedruckte Seitenzahl (sofern erkennbar).",
        "Relevanz-Score = interne Einschätzung der Passung zur Masterarbeit, keine Qualitätsbewertung.",
        "Status 'unsicher' bzw. Zitatprüfung 'not_found' => Angabe im Original-PDF prüfen.",
    ]:
        info.append([line])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_json() -> bytes:
    studies = []
    for s in services.list_studies():
        row = services.get_study_row(s["id"])
        studies.append({
            "id": s["id"], "filename": s["filename"], "status": s["status"], "imported_at": s["imported_at"],
            "page_count": s["page_count"], "text_quality": s["text_quality"], "relevance_score": s["relevance_score"],
            "user_tags": s["user_tags"], "notes": row.get("notes"), "chapter_override": s["chapter_override"],
            "analysis": db.loads(row.get("analysis_json"), None), "analysis_model": row.get("analysis_model"),
            "analyzed_at": row.get("analyzed_at"),
        })
    payload = {"exported_at": db.now_iso(), "app": "FCZ Thesis Research Hub", "studies": studies,
               "vocabulary": vocab()}
    return json.dumps(payload, ensure_ascii=False, indent=2).encode("utf-8")


def _bib_key(row: dict[str, Any]) -> str:
    authors = row.get("authors") or row.get("filename") or "study"
    first = re.split(r"[,;&]| and ", authors)[0].strip().split()
    name = re.sub(r"[^A-Za-z]", "", first[-1] if first else "study") or "study"
    return f"{name}{row.get('year') or 'nd'}_{row['id']}"


def to_bibtex() -> bytes:
    out = []
    for s in services.list_studies():
        if not (s["title"] or s["authors"]):
            continue
        authors = " and ".join(a.strip() for a in re.split(r";|,(?=\s*[A-ZÄÖÜ][a-zäöü]+\s)| & | and ", s["authors"] or "")
                               if a.strip())
        fields = {"title": s["title"], "author": authors, "year": s["year"], "journal": s["journal"], "doi": s["doi"]}
        body = ",\n".join(f"  {k} = {{{v}}}" for k, v in fields.items() if v)
        out.append(f"@article{{{_bib_key(s)},\n{body}\n}}")
    note = "% Automatisch extrahierte Metadaten – vor Verwendung prüfen.\n\n"
    return (note + "\n\n".join(out) + "\n").encode("utf-8")


def to_ris() -> bytes:
    out = []
    for s in services.list_studies():
        if not (s["title"] or s["authors"]):
            continue
        lines = ["TY  - JOUR"]
        if s["title"]:
            lines.append(f"TI  - {s['title']}")
        for a in re.split(r";|,(?=\s*[A-ZÄÖÜ][a-zäöü]+\s)| & | and ", s["authors"] or ""):
            if a.strip():
                lines.append(f"AU  - {a.strip()}")
        if s["year"]:
            lines.append(f"PY  - {s['year']}")
        if s["journal"]:
            lines.append(f"JO  - {s['journal']}")
        if s["doi"]:
            lines.append(f"DO  - {s['doi']}")
        lines.append("ER  - ")
        out.append("\n".join(lines))
    return ("\n\n".join(out) + "\n").encode("utf-8")


def _cite_str(cits: list[dict[str, Any]]) -> str:
    parts = []
    for c in cits or []:
        p = f", PDF-S. {c['page']}" if c.get("page") else ""
        if c.get("page_label"):
            p += f" (gedr. S. {c['page_label']})"
        parts.append(f"{c.get('ref', c.get('study_id'))}{p}")
    return f" [{'; '.join(parts)}]" if parts else " [ohne Beleg]"


def to_docx(chapter_keys: list[str] | None = None) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading("Kapitelnotizen – FCZ Thesis Research Hub", 0)
    doc.add_paragraph(f"Exportiert: {db.now_iso()}. Alle Belege verweisen auf physische PDF-Seiten "
                      "(gedruckte Seitenzahl in Klammern, sofern erkennbar). KI-Entwürfe sind gekennzeichnet und "
                      "müssen vor Verwendung geprüft werden.")
    for ch in CHAPTERS:
        if chapter_keys and ch["key"] not in chapter_keys:
            continue
        doc.add_heading(ch["title"], 1)
        note = db.get_conn().execute("SELECT content FROM user_notes WHERE scope='chapter' AND key=?",
                                     (ch["key"],)).fetchone()
        doc.add_heading("Eigene Notizen", 2)
        doc.add_paragraph(note["content"] if note and note["content"].strip() else "— keine —")

        studies = services.studies_for_chapter(ch["key"])
        doc.add_heading("Zugeordnete Studien (nach Relevanz)", 2)
        if not studies:
            doc.add_paragraph("— keine analysierten Studien zugeordnet —")
        for s in studies:
            row = services.get_study_row(s["id"])
            doc.add_paragraph(f"{services.short_ref(row)}: {s['title'] or s['filename']} "
                              f"(Relevanz {s['relevance_score']})", style="List Bullet")

        ai = services.latest_ai_output("theory", ch["key"])
        if ai:
            o = ai["output"]
            h = doc.add_heading("KI-Entwurf (ungeprüft – bitte verifizieren)", 2)
            for run in h.runs:
                run.font.color.rgb = RGBColor(0xB4, 0x53, 0x09)
            doc.add_paragraph(f"Erstellt {ai['created_at']} mit {ai['model']}. Nur auf Basis der importierten Studien.")
            for title, key in [("Zentrale Erkenntnisse", "key_insights"), ("Widersprüche", "contradictions"),
                               ("Argumentationsstruktur", "argument_structure")]:
                items = o.get(key) or []
                if items:
                    doc.add_heading(title, 3)
                    for it in items:
                        doc.add_paragraph(it.get("text", "") + _cite_str(it.get("citations")), style="List Bullet")
            if o.get("draft_paragraphs"):
                doc.add_heading("Textentwurf (KI)", 3)
                for it in o["draft_paragraphs"]:
                    doc.add_paragraph(it.get("text", "") + _cite_str(it.get("citations")))
            if o.get("missing_literature"):
                doc.add_heading("Fehlende Literatur (innerhalb der importierten Studien)", 3)
                for t in o["missing_literature"]:
                    doc.add_paragraph(t, style="List Bullet")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Backup / Restore
# ---------------------------------------------------------------------------

def create_backup() -> Path:
    stamp = db.now_iso().replace(":", "-").replace("+00-00", "Z")
    target = config.BACKUP_DIR / f"research_hub_backup_{stamp}.zip"
    with tempfile.TemporaryDirectory() as tmp:
        snap = Path(tmp) / "research_hub.sqlite3"
        src = db.get_conn()
        dst = sqlite3.connect(snap)
        with dst:
            src.backup(dst)
        dst.close()
        with zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as z:
            z.write(snap, "research_hub.sqlite3")
            for pdf in config.PDF_DIR.glob("*.pdf"):
                z.write(pdf, f"pdfs/{pdf.name}")
            z.writestr("backup_info.json", json.dumps({"created_at": db.now_iso(), "format": 1}))
    return target


def restore_backup(zip_bytes: bytes) -> dict[str, Any]:
    with tempfile.TemporaryDirectory() as tmp:
        zpath = Path(tmp) / "b.zip"
        zpath.write_bytes(zip_bytes)
        try:
            z = zipfile.ZipFile(zpath)
        except zipfile.BadZipFile as exc:
            raise ValueError("Keine gültige Backup-ZIP-Datei.") from exc
        with z:
            names = z.namelist()
            if "research_hub.sqlite3" not in names:
                raise ValueError("Backup enthält keine Datenbank (research_hub.sqlite3).")
            for n in names:
                if n.startswith("/") or ".." in Path(n).parts:
                    raise ValueError("Unsichere Pfade im Backup.")
            z.extractall(tmp)
        newdb = Path(tmp) / "research_hub.sqlite3"
        check = sqlite3.connect(newdb)
        try:
            n = check.execute("SELECT count(*) FROM studies").fetchone()[0]
        except sqlite3.DatabaseError as exc:
            raise ValueError("Datenbank im Backup ist ungültig.") from exc
        finally:
            check.close()
        # Sicherheitskopie des aktuellen Stands
        pre = create_backup()
        src = sqlite3.connect(newdb)
        with db._write_lock:
            dst = db.get_conn()
            src.backup(dst)
            src.close()
            dst.commit()
            for f in Path(tmp, "pdfs").glob("*.pdf") if Path(tmp, "pdfs").exists() else []:
                shutil.copy2(f, config.PDF_DIR / f.name)
    db.init_db()
    return {"restored_studies": n, "previous_state_backup": pre.name}
